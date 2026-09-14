"""Reproduce pandoc/skylighting SQL run-styling inside an existing .docx.

Harvests the token -> character-style map from the Source Code paragraphs already in the document,
so newly written SQL is highlighted exactly the way pandoc highlighted the rest of the file.
"""
import re
from collections import Counter, defaultdict

from docx.oxml.ns import qn

WORD = re.compile(r"[A-Za-z_][A-Za-z0-9_]*")
TOKENIZER = re.compile(r"""
    (?P<comment>--[^\n]*)
  | (?P<string>'(?:[^']|'')*')
  | (?P<number>\b\d+(?:\.\d+)?\b)
  | (?P<word>[A-Za-z_][A-Za-z0-9_]*)
  | (?P<space>[ \t]+)
  | (?P<other>.)
""", re.X)


def code_paragraphs(doc):
    out = []
    for p in doc.paragraphs:
        if p.style.name == "Source Code":
            out.append(p)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    if p.style.name == "Source Code":
                        out.append(p)
    return out


def harvest(doc):
    """word/operator -> most common character style, learned from the document itself."""
    votes = defaultdict(Counter)
    op_votes = defaultdict(Counter)
    for p in code_paragraphs(doc):
        for r in p.runs:
            style = r.style.name if r.style else None
            if style in (None, "Default Paragraph Font"):
                continue
            txt = r.text
            if style == "CommentTok" or style == "StringTok":
                continue
            for w in WORD.findall(txt):
                votes[w.upper()][style] += 1
            if style == "OperatorTok":
                for ch in txt.strip():
                    op_votes[ch][style] += 1
    word_style = {w: c.most_common(1)[0][0] for w, c in votes.items()}
    op_chars = set(op_votes)
    return word_style, op_chars


def segments(sql, word_style, op_chars):
    """-> list of (style, text) segments, adjacent same-style merged, '\\n' marking a line break."""
    segs = []
    for li, line in enumerate(sql.split("\n")):
        if li:
            segs.append(("__BR__", "\n"))
        for m in TOKENIZER.finditer(line):
            kind = m.lastgroup
            txt = m.group()
            if kind == "comment":
                style = "CommentTok"
            elif kind == "string":
                style = "StringTok"
            elif kind == "number":
                style = "FloatTok" if "." in txt else "DecValTok"
            elif kind == "word":
                style = word_style.get(txt.upper())
                if style is None:
                    # unseen identifier: skylighting styles a name followed by '(' as a function
                    after = line[m.end():m.end() + 1]
                    style = "FunctionTok" if after == "(" else "NormalTok"
            elif kind == "space":
                style = "NormalTok"
            else:
                style = "OperatorTok" if txt in op_chars else "NormalTok"
            if segs and segs[-1][0] == style:
                segs[-1] = (style, segs[-1][1] + txt)
            else:
                segs.append((style, txt))
    return segs


def _new_run(par, style, text, doc):
    r = par.add_run()
    if style == "__BR__":
        r._r.append(r._r.makeelement(qn("w:br"), {}))
        return r
    rPr = r._r.get_or_add_rPr()
    el = rPr.makeelement(qn("w:rStyle"), {qn("w:val"): style.replace(" ", "")})
    rPr.insert(0, el)
    r.text = text
    # keep leading/trailing spaces
    for t in r._r.findall(qn("w:t")):
        t.set(qn("xml:space"), "preserve")
    return r


def set_code(par, sql, doc, word_style, op_chars):
    for r in list(par.runs):
        r._r.getparent().remove(r._r)
    for style, text in segments(sql, word_style, op_chars):
        _new_run(par, style, text, doc)


def runs_signature(par):
    """(style, text) list for comparison, with <w:br/> as ('__BR__','\\n')."""
    out = []
    for r in par.runs:
        style = r.style.name if r.style else None
        if r._r.find(qn("w:br")) is not None and not r.text.strip():
            out.append(("__BR__", "\n"))
        else:
            out.append((("NormalTok" if style in (None, "Default Paragraph Font") else style).replace(" ", ""), r.text))
    return out
