"""Minimal GitHub-flavoured-markdown -> .docx converter (python-docx).

Stands in for pandoc, which can't run on this arm64 Mac (every local binary is x86-only).
Handles what docs/*.md actually use: ATX headings, pipe tables, fenced code blocks, bullet and
numbered lists (one nesting level), blockquotes, horizontal rules, and inline **bold** / *italic* /
`code` / [links](url).
"""
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

CODE_FONT = "Menlo"
CODE_BG = "F4F5F7"
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|\*[^*\s][^*]*?\*)", re.S)


def shade(par_or_cell, hexcolor):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexcolor)
    par_or_cell.get_or_add_pPr().append(el) if hasattr(par_or_cell, "get_or_add_pPr") else None


def shade_par(p, hexcolor):
    pPr = p._p.get_or_add_pPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexcolor)
    pPr.append(el)


def add_inline(par, text, base_bold=False, base_italic=False):
    """Split `text` into runs, honouring **bold**, *italic*, `code` and [label](url)."""
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            r = par.add_run(tok[2:-2].replace("\n", " "))
            r.bold = True
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            r = par.add_run(tok[1:-1])
            r.font.name = CODE_FONT
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0xB0, 0x30, 0x60)
        elif tok.startswith("[") and "](" in tok:
            label = tok[1:tok.index("](")]
            r = par.add_run(label)
            r.font.color.rgb = RGBColor(0x0B, 0x5C, 0xAD)
            r.underline = True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = par.add_run(tok[1:-1].replace("\n", " "))
            r.italic = True
        else:
            r = par.add_run(tok.replace("\n", " "))
        r.bold = r.bold or base_bold
        r.italic = r.italic or base_italic


def add_code_block(doc, lines, lang):
    for ln in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Inches(0.15)
        shade_par(p, CODE_BG)
        r = p.add_run(ln if ln.strip() else " ")
        r.font.name = CODE_FONT
        r.font.size = Pt(8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def split_row(line):
    cells = line.strip().strip("|").split("|")
    return [c.strip() for c in cells]


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=0, cols=len(header))
    t.style = "Table Grid"
    t.autofit = True
    hdr = t.add_row().cells
    for i, cell in enumerate(header):
        para = hdr[i].paragraphs[0]
        add_inline(para, cell, base_bold=True)
        el = OxmlElement("w:shd")
        el.set(qn("w:val"), "clear")
        el.set(qn("w:fill"), "EDEFF2")
        hdr[i]._tc.get_or_add_tcPr().append(el)
    for r in body:
        cells = t.add_row().cells
        for i in range(len(header)):
            add_inline(cells[i].paragraphs[0], r[i] if i < len(r) else "")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def convert(md_path, docx_path, title=None):
    src = open(md_path, encoding="utf-8").read().split("\n")
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)

    i, n = 0, len(src)
    while i < n:
        line = src[i]

        # fenced code
        if line.startswith("```"):
            lang = line[3:].strip()
            j = i + 1
            buf = []
            while j < n and not src[j].startswith("```"):
                buf.append(src[j])
                j += 1
            add_code_block(doc, buf, lang)
            i = j + 1
            continue

        # pipe table (needs the --- separator row underneath)
        if line.startswith("|") and i + 1 < n and re.match(r"^\|[\s:\-|]+\|$", src[i + 1].strip()):
            rows = [split_row(line)]
            j = i + 2
            while j < n and src[j].startswith("|"):
                rows.append(split_row(src[j]))
                j += 1
            add_table(doc, rows)
            i = j
            continue

        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if re.match(r"^(---+|\*\*\*+|___+)$", stripped):
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            bd = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "BBBBBB")
            bd.append(bottom)
            pPr.append(bd)
            i += 1
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            h = doc.add_heading(level=min(level, 4))
            add_inline(h, m.group(2))
            i += 1
            continue

        # blockquote (join continuation lines)
        if stripped.startswith(">"):
            buf = []
            while i < n and src[i].strip().startswith(">"):
                buf.append(re.sub(r"^\s*>\s?", "", src[i]))
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            shade_par(p, "F7F7F4")
            add_inline(p, " ".join(x.strip() for x in buf if x.strip()))
            continue

        # list item (bullet or numbered), with continuation lines
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if m:
            indent, marker, text = m.group(1), m.group(2), m.group(3)
            buf = [text]
            i += 1
            while i < n and src[i].strip() and not re.match(r"^(\s*)([-*]|\d+\.)\s+", src[i]) \
                    and not src[i].startswith(("#", "|", "```", ">")) \
                    and src[i].startswith((" ", "\t")):
                buf.append(src[i].strip())
                i += 1
            numbered = marker[0].isdigit()
            style = "List Number" if numbered else "List Bullet"
            if len(indent) >= 2:
                style += " 2"
            p = doc.add_paragraph(style=style)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, " ".join(buf))
            continue

        # plain paragraph (join wrapped lines)
        buf = [stripped]
        i += 1
        while i < n and src[i].strip() and not src[i].startswith(("#", "|", "```", ">", "---")) \
                and not re.match(r"^(\s*)([-*]|\d+\.)\s+", src[i]):
            buf.append(src[i].strip())
            i += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(buf))

    doc.save(docx_path)
    print("wrote %s (%d source lines)" % (docx_path, n))


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
