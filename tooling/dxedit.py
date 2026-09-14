"""Surgical .docx editing helpers that preserve pandoc's styling.

Text is written in light markdown (**bold**, *italic*, `code`) and rendered with the exact run
properties pandoc used: bold = <w:b/><w:bCs/>, code = rStyle VerbatimChar. Straight quotes are
converted to the typographic ones pandoc's --smart produces, so new prose matches the old.
"""
import copy
import re

from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

INLINE = re.compile(r"(\*\*`[^`]+`\*\*|\*\*[^*]+?\*\*|`[^`]+`|\*[^*\s][^*]*?\*)")


def body_items(doc):
    out = []
    for ch in doc.element.body.iterchildren():
        if ch.tag == qn("w:p"):
            out.append(Paragraph(ch, doc))
        elif ch.tag == qn("w:tbl"):
            out.append(Table(ch, doc))
    return out


def smart(text):
    """pandoc --smart: paired double quotes, apostrophes, ellipsis."""
    def dq(m):
        return "“" + m.group(1) + "”"
    text = re.sub(r'"([^"]*)"', dq, text)
    text = text.replace("...", "…")
    return text.replace("'", "’")


def _bold(run):
    rPr = run._r.get_or_add_rPr()
    for tag in ("w:b", "w:bCs"):
        rPr.append(rPr.makeelement(qn(tag), {}))


def _verbatim(run):
    rPr = run._r.get_or_add_rPr()
    rPr.insert(0, rPr.makeelement(qn("w:rStyle"), {qn("w:val"): "VerbatimChar"}))


def _italic(run):
    rPr = run._r.get_or_add_rPr()
    rPr.append(rPr.makeelement(qn("w:i"), {}))
    rPr.append(rPr.makeelement(qn("w:iCs"), {}))


def _add(par, text, bold=False, code=False, italic=False):
    if text == "":
        return
    r = par.add_run()
    if code:
        _verbatim(r)
    if bold:
        _bold(r)
    if italic:
        _italic(r)
    r.text = text if code else smart(text)
    for t in r._r.findall(qn("w:t")):
        t.set(qn("xml:space"), "preserve")


def clear_runs(par):
    for ch in list(par._p.iterchildren()):
        if ch.tag in (qn("w:r"), qn("w:hyperlink"), qn("w:bookmarkStart"), qn("w:bookmarkEnd")):
            par._p.remove(ch)


def write(par, md):
    """Replace a paragraph's content, keeping its paragraph style."""
    clear_runs(par)
    md = " ".join(md.split())
    for tok in INLINE.split(md):
        if not tok:
            continue
        if tok.startswith("**`") and tok.endswith("`**"):
            _add(par, tok[3:-3], bold=True, code=True)
        elif tok.startswith("**") and tok.endswith("**"):
            _add(par, tok[2:-2], bold=True)
        elif tok.startswith("`") and tok.endswith("`"):
            _add(par, tok[1:-1], code=True)
        elif tok.startswith("*") and tok.endswith("*"):
            _add(par, tok[1:-1], italic=True)
        else:
            _add(par, tok)
    return par


def clone_para(template, after_el, md=None, style=None):
    """Deep-copy a paragraph element (keeping its style/numbering) and insert it after after_el."""
    new = copy.deepcopy(template._p)
    after_el.addnext(new)
    par = Paragraph(new, template._parent)
    if style:
        par.style = style
    if md is not None:
        write(par, md)
    return par


def cell_write(cell, md):
    par = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    write(par, md)


def set_row(table, ri, values):
    cells = table.rows[ri].cells
    for i, v in enumerate(values):
        if i < len(cells):
            cell_write(cells[i], v)


def insert_row_after(table, ri, values):
    """Clone row `ri` (preserving its cell formatting) and write `values` into the copy."""
    src = table.rows[ri]._tr
    new = copy.deepcopy(src)
    src.addnext(new)
    idx = ri + 1
    set_row(table, idx, values)
    return table.rows[idx]


def clone_table_after(table, after_el):
    new = copy.deepcopy(table._tbl)
    after_el.addnext(new)
    return Table(new, table._parent)


def find_item(items, pred):
    for i, o in enumerate(items):
        if isinstance(o, Paragraph) and pred(o):
            return i
    raise LookupError("not found")


def add_bookmark(par, name):
    """Re-attach a pandoc-style heading anchor after a heading's text has been rewritten."""
    body = par._p.getparent()
    while body.tag != qn("w:body"):
        body = body.getparent()
    used = [int(e.get(qn("w:id"))) for e in body.iter(qn("w:bookmarkStart")) if e.get(qn("w:id")) is not None]
    bid = str(max(used) + 1 if used else 1)
    start = par._p.makeelement(qn("w:bookmarkStart"), {qn("w:id"): bid, qn("w:name"): name})
    end = par._p.makeelement(qn("w:bookmarkEnd"), {qn("w:id"): bid})
    par._p.insert(1, start)
    par._p.insert(2, end)
    return name
